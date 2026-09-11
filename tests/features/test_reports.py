"""Operations over two days and two branches add up in every report."""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import parse_qs, urlparse

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status

from apps.companies.models.operation_model import CarOperation
from apps.companies.tests.api.v1.car_operation.helpers import (
    operation_download_url,
    operation_export_url,
)
from apps.companies.tests.helpers import set_balance
from apps.stations.tests.helpers import operations_url, reports_url

from .helpers import ALL_DAYS, company_home_url, fuel, sign_in

pytestmark = [pytest.mark.django_db, pytest.mark.feature]

TOTAL_LABEL = "الاجمالي"


def petrol_row(report):
    (row,) = report.data["operations"]
    return {key: row[key] for key in ("total_balance", "count", "amount")}


def export_totals(client, **params):
    """Export, download the file and return each car's total cost cell."""
    exported = client.get(operation_export_url(), params)
    if exported.status_code != status.HTTP_200_OK:
        raise RuntimeError(f"export returned {exported.status_code}: {exported.data}")
    filename = parse_qs(urlparse(exported.data["download_url"]).query)["file"][0]
    downloaded = client.get(operation_download_url(), {"file": filename})
    sheet = load_workbook(BytesIO(b"".join(downloaded.streaming_content))).active
    return sorted(
        row[0] for row in sheet.iter_rows(values_only=True) if row[-1] == TOTAL_LABEL
    )


class TestReports:
    @pytest.fixture(autouse=True)
    def setup(
        self,
        fees,
        fuelable_car,
        car_factory,
        car_operation_factory,
        company_driver,
        company_owner,
        company_branch_manager,
        second_company_branch,
        branch,
        second_station_branch,
        station_owner,
        branch_manager,
        station_worker,
        second_station_worker,
    ):
        second_company_branch.fees = Decimal("10.00")
        second_company_branch.save(update_fields=["fees"])
        second_station_branch.fees = Decimal("0.50")
        second_station_branch.save(update_fields=["fees"])
        self.car_a = fuelable_car
        self.car_b = car_factory(
            branch=second_company_branch, fuel_allowed_days=ALL_DAYS, last_meter=10000
        )
        for item in (self.car_a, self.car_b, branch, second_station_branch):
            set_balance(item, "1000.00")

        # yesterday car A took 5 L at station branch 1
        yesterday = timezone.localtime() - timedelta(days=1)
        self.yesterday_op = car_operation_factory(
            car=self.car_a,
            status=CarOperation.OperationStatus.COMPLETED,
            amount=Decimal("5.00"),
            cost=Decimal("50.00"),
            company_cost=Decimal("55.00"),
            station_cost=Decimal("52.50"),
            profits=Decimal("2.50"),
            car_first_meter=Decimal("9950.00"),
        )
        CarOperation.objects.filter(id=self.yesterday_op.id).update(
            created=yesterday,
            modified=yesterday,
            start_time=yesterday,
            end_time=yesterday,
        )
        # today car A takes 20 L at branch 1, car B 10 L at branch 2
        self.op_a, _ = fuel(
            sign_in("station", station_worker),
            company_driver,
            self.car_a,
            amount="20",
            meter="10100",
        )
        self.op_b, _ = fuel(
            sign_in("station", second_station_worker),
            company_driver,
            self.car_b,
            amount="10",
            meter="10100",
        )
        self.today = timezone.localdate().isoformat()
        self.yesterday = yesterday.date().isoformat()
        self.company_owner = sign_in("company", company_owner)
        self.company_manager = sign_in("company", company_branch_manager)
        self.station_owner = sign_in("station", station_owner)
        self.station_manager = sign_in("station", branch_manager)
        self.second_worker = sign_in("station", second_station_worker)

    def test_station_reports_add_up_per_day_success(self):
        today = self.station_owner.get(
            reports_url(date_from=self.today, date_to=self.today)
        )
        both_days = self.station_owner.get(
            reports_url(date_from=self.yesterday, date_to=self.today)
        )
        yesterday = self.station_owner.get(
            reports_url(date_from=self.yesterday, date_to=self.yesterday)
        )
        listed = self.station_owner.get(operations_url())

        # 20 L and 10 L at 10.50/L for the station
        assert petrol_row(today) == {
            "total_balance": "315.00",
            "count": 2,
            "amount": "30.00",
        }
        assert petrol_row(both_days) == {
            "total_balance": "367.50",
            "count": 3,
            "amount": "35.00",
        }
        assert petrol_row(yesterday)["total_balance"] == "52.50"
        assert listed.data["petrol_balance"] == Decimal("367.50")
        assert listed.data["count"] == 3

    def test_station_reports_are_scoped_to_the_viewer_success(self):
        params = {"date_from": self.today, "date_to": self.today}

        manager = self.station_manager.get(reports_url(**params))
        worker = self.second_worker.get(reports_url(**params))

        # the manager runs branch 1 (car A), the second worker branch 2 (car B)
        assert petrol_row(manager)["total_balance"] == "210.00"
        assert petrol_row(worker)["total_balance"] == "105.00"

    def test_company_home_lists_the_latest_operations_success(self):
        home = self.company_owner.get(company_home_url())

        assert [item["id"] for item in home.data["car_operations"]] == [
            self.op_b,
            self.op_a,
            self.yesterday_op.id,
        ]

    def test_export_matches_what_the_operations_charged_success(self):
        owner_today = export_totals(
            self.company_owner, date_from=self.today, date_to=self.today
        )
        owner_car_a = export_totals(
            self.company_owner,
            date_from=self.yesterday,
            date_to=self.today,
            car=self.car_a.id,
        )
        manager_today = export_totals(
            self.company_manager, date_from=self.today, date_to=self.today
        )

        assert owner_today == ["110.00 جنيه", "220.00 جنيه"]
        # 220 today + 55 yesterday
        assert owner_car_a == ["275.00 جنيه"]
        # the manager runs branch A only
        assert manager_today == ["220.00 جنيه"]

    def test_home_counts_cars_with_equal_balances_once_fail(self):
        """Open issue: `cars_balance` is `Sum(..., distinct=True)`, which sums
        distinct values, so two cars holding 780 each report 780, not 1560."""
        set_balance(self.car_b, "780.00")

        home = self.company_owner.get(company_home_url())

        assert home.data["cars_balance"] == Decimal("780.00")
